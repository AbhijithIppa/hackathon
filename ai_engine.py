"""StatGrow AI competency engine — simple functions with LLM or mock fallback."""

from __future__ import annotations

import json
import os
import random
import re
from typing import Any

from demo_data import (
    ASSISTANT_RESPONSES,
    COMPETENCIES,
    DEFAULT_PROFILE,
    IGOT_COURSES,
    LEARNING_MATERIALS,
    LEARNING_MODULES,
    QUIZ_BANK,
)

# Reason: Keep LLM optional so hackathon demos never break without a key.
def _api_key() -> str | None:
    return os.environ.get("OPENAI_API_KEY") or os.environ.get("STATGROW_LLM_API_KEY")


def _llm_available() -> bool:
    return bool(_api_key())


def _call_llm(system: str, user: str) -> str | None:
    """Best-effort OpenAI-compatible chat completion. Returns None on failure."""
    if not _llm_available():
        return None
    try:
        from openai import OpenAI

        client = OpenAI(api_key=_api_key())
        model = os.environ.get("OPENAI_MODEL", "gpt-4o-mini")
        resp = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
            temperature=0.4,
        )
        return (resp.choices[0].message.content or "").strip()
    except Exception:
        return None


def _parse_json_block(text: str) -> Any | None:
    if not text:
        return None
    text = text.strip()
    fence = re.search(r"```(?:json)?\s*([\s\S]*?)```", text)
    if fence:
        text = fence.group(1).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("["), text.rfind("]")
        if start != -1 and end > start:
            try:
                return json.loads(text[start : end + 1])
            except json.JSONDecodeError:
                pass
        start, end = text.find("{"), text.rfind("}")
        if start != -1 and end > start:
            try:
                return json.loads(text[start : end + 1])
            except json.JSONDecodeError:
                pass
    return None


def _bank_for(competency: str) -> list[dict]:
    if competency in QUIZ_BANK and QUIZ_BANK[competency]:
        return list(QUIZ_BANK[competency])
    # Fallback: mix from related banks
    pooled: list[dict] = []
    for qs in QUIZ_BANK.values():
        pooled.extend(qs)
    return pooled


def _normalize_question(q: dict, default_competency: str, difficulty: str) -> dict:
    options = list(q.get("options") or [])
    while len(options) < 4:
        options.append(f"Option {chr(65 + len(options))}")
    options = options[:4]
    correct = q.get("correct", 0)
    if isinstance(correct, str):
        letter = correct.strip().upper()[:1]
        if letter in "ABCD":
            correct = ord(letter) - ord("A")
        else:
            try:
                correct = options.index(correct)
            except ValueError:
                correct = 0
    correct = int(correct) % 4
    return {
        "question": q.get("question", "Untitled question"),
        "options": options,
        "correct": correct,
        "explanation": q.get(
            "explanation",
            "Review the related concept in your learning path.",
        ),
        "competency": q.get("competency", default_competency),
        "difficulty": q.get("difficulty", difficulty),
        "type": q.get("type", "MCQ"),
    }


def extract_topics(text: str) -> list[str]:
    """Identify topics from learning material text."""
    llm = _call_llm(
        "You extract topics for India's Official Statistical System. Return JSON array of 4-8 short topic strings.",
        f"Extract topics from:\n{text[:6000]}",
    )
    parsed = _parse_json_block(llm or "")
    if isinstance(parsed, list) and parsed:
        return [str(t) for t in parsed[:8]]

    lower = text.lower()
    topic_map = {
        "stratified": "Stratified Sampling",
        "cluster": "Cluster Sampling",
        "probability": "Probability Sampling",
        "sampling error": "Sampling Error",
        "census": "Census Methodology",
        "nss": "NSS Survey Design",
        "data quality": "Data Quality",
        "visualization": "Data Visualization",
        "governance": "Data Governance",
        "cpi": "Price Statistics",
        "national account": "National Accounts",
        "questionnaire": "Questionnaire Design",
        "non-response": "Non-response Treatment",
    }
    found = [label for key, label in topic_map.items() if key in lower]
    if not found:
        found = ["Probability Sampling", "Stratified Sampling", "Cluster Sampling", "Sampling Error"]
    return found[:8]


def extract_competencies(text: str, topics: list[str] | None = None) -> list[str]:
    """Map text/topics to competency labels."""
    blob = (text or "") + " " + " ".join(topics or [])
    llm = _call_llm(
        f"Map content to competencies from this list only: {COMPETENCIES}. Return JSON array of 2-5 names.",
        f"Content:\n{blob[:4000]}",
    )
    parsed = _parse_json_block(llm or "")
    if isinstance(parsed, list):
        mapped = [c for c in parsed if c in COMPETENCIES]
        if mapped:
            return mapped[:5]

    lower = blob.lower()
    rules = [
        (["sampl", "strata", "cluster", "psu"], "Sampling Techniques"),
        (["survey", "questionnaire", "enumerat"], "Survey Methodology"),
        (["visual", "chart", "dashboard"], "Data Visualization"),
        (["quality", "dqaf", "timeliness"], "Data Quality"),
        (["governance", "steward", "confidential"], "Data Governance"),
        (["census"], "Census & Surveys"),
        (["economic", "cpi", "gdp", "iip"], "Economic Statistics"),
        (["social", "labour", "education", "health"], "Social Statistics"),
        (["official statistic", "mospi", "nso"], "Official Statistics"),
        (["machine learning", "ai ", " ml"], "AI & Machine Learning"),
        (["analysis", "regression", "inference"], "Data Analysis"),
        (["python", "r studio", "stata", "software"], "Statistical Software"),
        (["collect", "capi", "field"], "Data Collection"),
        (["statistical method", "hypothesis", "variance"], "Statistical Methods"),
    ]
    hits = []
    for keys, comp in rules:
        if any(k in lower for k in keys) and comp not in hits:
            hits.append(comp)
    return hits[:5] or ["Sampling Techniques", "Survey Methodology", "Statistical Methods"]


def analyze_learning_material(
    text: str,
    filename: str = "Uploaded Document",
) -> dict[str, Any]:
    """Full document analysis pipeline for Knowledge Studio."""
    topics = extract_topics(text)
    competencies = extract_competencies(text, topics)
    difficulty = "Intermediate"
    if any(w in text.lower() for w in ["advanced", "design effect", "calibration"]):
        difficulty = "Advanced"
    elif any(w in text.lower() for w in ["introduction", "basics", "overview"]):
        difficulty = "Beginner"

    questions = generate_mcqs(
        topic=topics[0] if topics else "Official Statistics",
        n=10,
        difficulty="Medium",
        question_type="MCQ",
        competency=competencies[0] if competencies else "Sampling Techniques",
        source_material=text[:2000],
    )
    recommendations = [
        f"Deepen practice on {topics[0]}" if topics else "Review core concepts",
        f"Link assessment to {competencies[0]}" if competencies else "Take a competency assessment",
        "Generate a quiz from this material in AI Quiz Generator",
    ]
    return {
        "document": filename,
        "topics": topics,
        "competencies": competencies,
        "difficulty": difficulty,
        "questions": questions,
        "question_count": len(questions),
        "recommendations": recommendations,
        "summary": (
            f"Analyzed '{filename}': mapped to {', '.join(competencies[:3])} "
            f"with focus on {', '.join(topics[:3])}."
        ),
    }


def generate_mcqs(
    topic: str,
    n: int = 5,
    difficulty: str = "Medium",
    question_type: str = "MCQ",
    competency: str = "Sampling Techniques",
    source_material: str = "",
) -> list[dict]:
    """Generate assessment/quiz questions."""
    n = max(1, min(int(n), 15))
    system = (
        "You create high-quality assessment questions for India's Official Statistical System. "
        "Return JSON array of objects with keys: question, options (4 strings), correct (0-3 index), "
        "explanation, competency, difficulty, type."
    )
    user = (
        f"Topic: {topic}\nCompetency: {competency}\nDifficulty: {difficulty}\n"
        f"Type: {question_type}\nCount: {n}\n"
        f"Source context:\n{source_material[:3000]}"
    )
    llm = _call_llm(system, user)
    parsed = _parse_json_block(llm or "")
    if isinstance(parsed, list) and parsed:
        return [_normalize_question(q, competency, difficulty) for q in parsed[:n]]

    bank = _bank_for(competency)
    # Prefer matching difficulty / type when possible
    preferred = [
        q
        for q in bank
        if q.get("difficulty", "").lower() == difficulty.lower()
        or q.get("type", "").lower() == question_type.lower()
        or topic.lower().split()[0] in q.get("question", "").lower()
    ]
    pool = preferred or bank
    random.shuffle(pool)
    selected = pool[:n]
    # If bank is short, cycle with slight variation labels
    while len(selected) < n and bank:
        base = dict(random.choice(bank))
        base["question"] = base["question"] + f" (Practice variant {len(selected) + 1})"
        selected.append(base)
    return [_normalize_question(q, competency, difficulty) for q in selected[:n]]


def generate_assessment(competency: str, n: int = 8) -> list[dict]:
    """Generate a competency assessment (5–10 questions)."""
    n = max(5, min(int(n), 10))
    return generate_mcqs(
        topic=competency,
        n=n,
        difficulty="Medium",
        question_type="MCQ",
        competency=competency,
        source_material=f"Assessment for competency: {competency} in India's official statistics context.",
    )


def evaluate_assessment(
    questions: list[dict],
    answers: list[int | None],
) -> dict[str, Any]:
    """Score an assessment and produce strengths / gaps."""
    total = len(questions)
    correct_flags = []
    by_comp: dict[str, list[bool]] = {}
    weak_topics: list[str] = []

    for i, q in enumerate(questions):
        ans = answers[i] if i < len(answers) else None
        ok = ans is not None and int(ans) == int(q.get("correct", -1))
        correct_flags.append(ok)
        comp = q.get("competency", "General")
        by_comp.setdefault(comp, []).append(ok)
        if not ok:
            # Derive a short gap label from question keywords
            text = q.get("question", "")
            for key in [
                "stratified",
                "cluster",
                "sampling error",
                "sample size",
                "design weight",
                "non-response",
                "visualization",
                "quality",
                "census",
            ]:
                if key in text.lower():
                    label = key.title()
                    if label not in weak_topics:
                        weak_topics.append(label)

    score_pct = round(100 * sum(correct_flags) / total) if total else 0
    strengths = [c for c, flags in by_comp.items() if flags and sum(flags) / len(flags) >= 0.7]
    weak_areas = [c for c, flags in by_comp.items() if flags and sum(flags) / len(flags) < 0.6]
    if not weak_topics:
        weak_topics = weak_areas[:3] or ["Core concepts in selected competency"]

    comp_scores = {
        c: round(100 * sum(flags) / len(flags)) for c, flags in by_comp.items() if flags
    }
    recommended = []
    for w in weak_areas or weak_topics:
        recommended.append(f"Review: {w}")
    if not recommended:
        recommended = ["Advance to next difficulty level", "Try scenario-based questions"]

    return {
        "score": sum(correct_flags),
        "total": total,
        "score_pct": score_pct,
        "correct": sum(correct_flags),
        "incorrect": total - sum(correct_flags),
        "strengths": strengths or ["Consistent attempt across questions"],
        "weak_areas": weak_areas or (["Needs reinforcement"] if score_pct < 80 else []),
        "knowledge_gaps": weak_topics[:5],
        "recommended_topics": recommended[:5],
        "competency_scores": comp_scores,
        "details": correct_flags,
    }


def analyze_competency_gap(profile: dict[str, int], threshold: int = 60) -> dict[str, Any]:
    """Identify priority gaps from a competency profile."""
    ranked = sorted(profile.items(), key=lambda x: x[1])
    gaps = [(name, score) for name, score in ranked if score < threshold]
    priority = gaps[0] if gaps else ranked[0]
    return {
        "priority": {"name": priority[0], "score": priority[1]},
        "gaps": [{"name": n, "score": s} for n, s in gaps[:5]],
        "ranked_asc": [{"name": n, "score": s} for n, s in ranked],
    }


def level_from_score(score: int) -> str:
    if score < 50:
        return "Beginner"
    if score < 75:
        return "Intermediate"
    return "Advanced"


def target_level(score: int) -> str:
    current = level_from_score(score)
    return {"Beginner": "Intermediate", "Intermediate": "Advanced", "Advanced": "Expert"}.get(
        current, "Advanced"
    )


def generate_learning_path(profile: dict[str, int], top_n: int = 3) -> list[dict]:
    """Build personalized learning path from weakest competencies."""
    gaps = analyze_competency_gap(profile)["ranked_asc"][:top_n]
    path = []
    for item in gaps:
        name, score = item["name"], item["score"]
        modules = LEARNING_MODULES.get(name)
        if not modules:
            modules = [
                {
                    "title": f"Foundations of {name}",
                    "overview": f"Core concepts and practice for {name}.",
                    "objectives": [f"Explain key ideas in {name}", "Apply concepts in official statistics"],
                    "duration": "45 min",
                    "resources": [f"{name} reading pack (demo)"],
                },
                {
                    "title": f"Applied {name}",
                    "overview": f"Case-based application for {name}.",
                    "objectives": ["Solve applied scenarios", "Self-check with a short quiz"],
                    "duration": "50 min",
                    "resources": [f"{name} case studies (demo)"],
                },
            ]
        path.append(
            {
                "competency": name,
                "current_level": level_from_score(score),
                "target_level": target_level(score),
                "progress": max(10, min(score, 95)),
                "score": score,
                "modules": modules,
            }
        )
    return path


def generate_course_recommendations(profile: dict[str, int], top_n: int = 3) -> list[dict]:
    """Mock iGOT recommendations from competency gaps.

    Structured so a future real iGOT API can replace this function body.
    """
    gaps = analyze_competency_gap(profile)["ranked_asc"]
    recs = []
    used = set()
    for g in gaps:
        for course in IGOT_COURSES:
            if course["competency"] == g["name"] and course["id"] not in used:
                recs.append(
                    {
                        **course,
                        "gap_score": g["score"],
                        "why": (
                            f"Your assessment indicates a competency gap in {g['name']} "
                            f"(current score: {g['score']}%)."
                        ),
                    }
                )
                used.add(course["id"])
                break
        if len(recs) >= top_n:
            break
    return recs


def explain_answer(question: dict, user_answer: int | None) -> str:
    """Explain correct answer and why a user response may be wrong."""
    correct_idx = int(question.get("correct", 0))
    options = question.get("options", [])
    correct_text = options[correct_idx] if correct_idx < len(options) else ""
    base = question.get("explanation", "")

    llm = _call_llm(
        "Explain quiz answers clearly for statistical officers in India. Keep under 120 words.",
        f"Question: {question.get('question')}\nOptions: {options}\n"
        f"Correct index: {correct_idx}\nUser answer index: {user_answer}\nExplanation: {base}",
    )
    if llm:
        return llm

    if user_answer is None:
        return f"Correct answer: **{correct_text}**.\n\n{base}"
    if int(user_answer) == correct_idx:
        return f"Correct — **{correct_text}**.\n\n{base}"
    chosen = options[user_answer] if 0 <= int(user_answer) < len(options) else "—"
    return (
        f"You selected **{chosen}**. The correct answer is **{correct_text}**.\n\n{base}\n\n"
        "Tip: contrast definitions carefully (e.g., stratified vs cluster sampling)."
    )


def update_profile_from_result(
    profile: dict[str, int],
    result: dict[str, Any],
    primary_competency: str | None = None,
    boost_cap: int = 25,
) -> dict[str, int]:
    """Blend assessment outcome into session competency profile (demo-friendly)."""
    updated = dict(profile)
    # Primary competency gets a clear before/after movement for demos
    if primary_competency:
        current = updated.get(primary_competency, 50)
        # Map score_pct into an uplift toward mastery for demo narrative
        target = max(current, min(95, int(0.45 * current + 0.55 * result["score_pct"])))
        # Ensure visible improvement when user does reasonably well
        if result["score_pct"] >= 60:
            target = max(target, min(95, current + max(8, result["score_pct"] // 5)))
        uplift = min(boost_cap, max(0, target - current))
        updated[primary_competency] = current + uplift

    for comp, pct in (result.get("competency_scores") or {}).items():
        cur = updated.get(comp, 50)
        blended = int(0.6 * cur + 0.4 * pct)
        if pct >= 70:
            blended = max(blended, min(95, cur + 5))
        updated[comp] = max(5, min(95, blended))
    return updated


def assistant_reply(message: str, context: dict | None = None) -> str:
    """StatGrow AI Assistant response."""
    ctx = context or {}
    llm = _call_llm(
        "You are StatGrow AI Assistant for India's Official Statistical System. "
        "Be concise, practical, and supportive. Help with concepts, gaps, and practice.",
        f"User message: {message}\nContext: {json.dumps(ctx)[:2000]}",
    )
    if llm:
        return llm

    m = message.lower()
    if "stratif" in m or "simply" in m:
        return ASSISTANT_RESPONSES["stratified"]
    if "example" in m:
        return ASSISTANT_RESPONSES["example"]
    if "test my" in m or "test me" in m:
        return ASSISTANT_RESPONSES["test"]
    if "mcq" in m or "5" in m and "question" in m:
        return ASSISTANT_RESPONSES["mcq"]
    if "wrong" in m or "incorrect" in m:
        return ASSISTANT_RESPONSES["wrong"]
    if "gap" in m:
        gap = ctx.get("priority_gap") or analyze_competency_gap(
            ctx.get("profile") or DEFAULT_PROFILE
        )["priority"]
        return (
            f"Your priority competency gap is **{gap['name']}** at **{gap['score']}%**.\n\n"
            "Recommended next step: open **My Learning Path**, complete the first module, "
            "then take a practice quiz in **AI Quiz Generator**."
        )
    return ASSISTANT_RESPONSES["default"]


def demo_material_text(material_id: str = "mat_sampling") -> str:
    """Return synthetic document text for demo processing without uploads."""
    for m in LEARNING_MATERIALS:
        if m["id"] == material_id:
            topics = ", ".join(m["topics"])
            comps = ", ".join(m["competencies"])
            return (
                f"# {m['title']}\n\n{m['summary']}\n\n"
                f"Topics covered: {topics}.\n"
                f"Related competencies: {comps}.\n\n"
                "This material covers probability sampling, stratified sampling, cluster sampling, "
                "sampling error, design weights, and multi-stage designs used in NSS and Census "
                "operations within India's Official Statistical System."
            )
    return (
        "Survey Sampling Fundamentals covering probability sampling, stratified sampling, "
        "cluster sampling and sampling error for official statistics."
    )


def extract_text_from_upload(name: str, raw: bytes) -> str:
    """Best-effort local text extraction for PDF/DOCX/PPTX/TXT."""
    lower = name.lower()
    try:
        if lower.endswith(".txt"):
            return raw.decode("utf-8", errors="ignore")
        if lower.endswith(".pdf"):
            import io

            from PyPDF2 import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            parts = []
            for page in reader.pages[:20]:
                parts.append(page.extract_text() or "")
            text = "\n".join(parts).strip()
            return text or demo_material_text()
        if lower.endswith(".docx"):
            import io

            from docx import Document

            doc = Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs if p.text).strip()
            return text or demo_material_text()
        if lower.endswith(".pptx"):
            import io

            from pptx import Presentation

            prs = Presentation(io.BytesIO(raw))
            chunks = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        chunks.append(shape.text)
            text = "\n".join(chunks).strip()
            return text or demo_material_text()
    except Exception:
        pass
    # Reason: Always return usable text so Knowledge Studio demo continues.
    return demo_material_text() + f"\n\n[Processed placeholder for: {name}]"


def overall_competency(profile: dict[str, int]) -> int:
    if not profile:
        return 0
    return round(sum(profile.values()) / len(profile))
